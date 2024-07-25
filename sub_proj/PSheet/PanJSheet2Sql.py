# %%
import numpy as np
import tqdm
import re
import datetime
import sys
from typing import List, Tuple, Dict, Union, Optional, Any

from openpyxl import load_workbook
from openpyxl.cell.cell import Cell
from openpyxl.styles.colors import Color

from ExcelColor import theme_and_tint_to_rgb

# %%
input_path = "Z:\\Proj\\Jyutdict\\泛粵字表\\本体\\泛粵字表 240723.xlsx"
output_dir = "Z:\\Proj\\Jyutdict\\泛粵字表\\Automatic\\"

time_str = datetime.datetime.now().strftime('%Y%m%d')
regex_pure_alphabet = re.compile("^[a-z'0-9?/①-⑨_^*]+$")
def str_format_enter_for_12(x: str) -> str:
    if len(regex_pure_alphabet.findall(x))>0: return x
    x = x.replace("；", ";")
    x = x.replace("，", ",")
    x = x.replace("：", ":")
    x = x.replace(", ", ",")
    x = x.replace("; ", ";")
    x = x.replace(",\n", ",")
    x = x.replace(";\n", ";")
    x = x.replace(",", ", ")
    x = x.replace(";", "; ")
    x = x.replace("  ", " ")
    x = x.replace(" :", ": ")
    x = x.replace("\n", "; ")
    x = x.replace("  ", " ")
    return x.strip()
def str_format_enter_for_m2(x: str) -> str:
    if "需要例句" in x:
        x = x.replace("：需要例句", "")
        x = x.replace("，需要例句", "")
    # x = x.replace("⚠意味不明", "{？}").strip()
    if x[-1]=="$": x = x[:-1]
    return x


HEADER_INFO_COL_LENGTH = []
HEADER_INFO_COL_NAME  = ['綜', '釋', '繁', ]
HEADER_INFO_MARK      = [0, 0, 0]
HEADER_INFO_FULL_NAME = [['綜合音', ''], ['簡釋', ''], ['字', '']]
HEADER_INFO_COLOR     = ['', '', '']
HEADER_INFO_NOTE      = ['', '', '']


regex_block_email = re.compile("@[a-zA-Z0-9\\.]+?@[a-zA-Z0-9]+?\\.[a-zA-Z]+\\b")
regex_block_name = re.compile("\n\t-.+?(\n|$)")
def regex_block_info(s: str) -> str:
    if not s: return ""
    x = s.strip()
    if "@" in x: x = regex_block_email.sub("@[Anony]", x)
    if "-" in x: x = regex_block_name.sub("\t-[Anony]\n", x)
    return x

content_block_pair = {
    re.compile("\\\\xa0"                                         ): r"\\t",
    re.compile("@[a-zA-Z0-9.]+?@[a-zA-Z0-9]+?\\\\.[a-zA-Z]+\\\\b"): r"[Anony]",
    re.compile("\\\\n_已指派給[^_]+?_"                           ): r"",
    re.compile("\\\\n[-_]{10,}\\\\n"                             ): r"\\n",
    re.compile("\\\\n[-_]{10,}"                                  ): r"",
    re.compile("\\\\n\\\\t-[^']+?(\\\\n|')"                      ): r"\\n\\t-Anony\1",
    re.compile("('|\\\\n)[^\\\\:]+?:\\\\n"                       ): r"\1[Anony]:\\t",
    re.compile("\\\\n_Marked as resolved_\\\\n\\\\t-[Anony]"     ): r"",
    re.compile("\\\\n_Re-opened_\\\\n\\\\t-[Anony]"              ): r"",
    re.compile("\\\\n-------------------------$"                 ): r"",
}
def regex_block_note(s: str) -> str:
    if not s: return ""
    x = s.strip()
    for k, v in content_block_pair.items():
        x = k.sub(v, x)
    return x

def format_locale_record(row: int) -> Tuple[List[str], List[int], bool]:
    sheet_row = main_sheet[row]
    assert isinstance(sheet_row, tuple)
    valid = True
    result: List[str] = [""] * (HEADER_INFO_COL_COUNT+1)
    col_length = []
    locate_written_count = 0
    locate_owned_count = 0
    col_count = len(sheet_row)
    for n, item in enumerate(sheet_row):
        if item.value:
            value = str(item.value).strip()
            if '"' in value and "'" in value:
                value = value.replace('"', "'")
            # if "需要例句" in value:
            #     print(value, HEADER_INFO_MARK[n], type(HEADER_INFO_MARK[n]))
            match HEADER_INFO_MARK[n]:
                case 1:
                    value = str_format_enter_for_12(value)
                    locate_written_count += 1
                    if value != "_":
                        locate_owned_count += 1
                case 2:
                    value = str_format_enter_for_12(value)
                case -2:
                    if "⚠" in value:
                        valid = False
                        break
                    value = str_format_enter_for_m2(value)
            
            if "'" not in value:
                result[n] = "'%s'" % value
            else:
                result[n] = '"%s"' % value
        else:
            result[n] = "''"
    if locate_written_count<3 and locate_owned_count<2:
        valid = False
    
    if not valid:
        return result, [], False
    
    comment_filtered = [(n, item.comment.text) for n, item in enumerate(sheet_row) if item.comment]
    result_note = str({
        HEADER_INFO_COL_NAME[n]: 
        regex_block_info(comment) 
        for n, comment in comment_filtered
    }).replace("\"", "\\\"")
    result[col_count] = f'"{regex_block_note(result_note)}"'
    
    col_length = [len(item)-2 for item in result]
    return result, col_length, True
    
# %%
print("0___讀取字表___")
PanCSheet = load_workbook(filename = input_path, data_only=True, keep_vba=False)
HEADER_ROW_INDEX = 1
HEADERMARK_ROW_INDEX = 4
HEADERNAME_ROW_INDEX = 6

main_sheet = PanCSheet["主表-睇真尐註解"]
for n in tqdm.trange(main_sheet.max_column, 0, -1):
    if main_sheet.cell(row=HEADERMARK_ROW_INDEX, column=n).value == "x":
        main_sheet.delete_cols(n)
headers      = main_sheet[HEADER_ROW_INDEX]
headers_mark = main_sheet[HEADERMARK_ROW_INDEX]
headers_name = main_sheet[HEADERNAME_ROW_INDEX]
assert isinstance(headers, tuple) and isinstance(headers_mark, tuple) and isinstance(headers_name, tuple)

sql_ifaamjyut_header = """CREATE TABLE `IFaamjyut` (
  `id` int(11) NOT NULL,
  `col` varchar(10) CHARACTER SET utf8mb4 COLLATE utf8mb4_bin DEFAULT NULL,
  `kind` int(11) DEFAULT '0',
  `fullname` varchar(20) CHARACTER SET utf8mb4 COLLATE utf8mb4_bin DEFAULT NULL,
  `fullname_note` varchar(20) CHARACTER SET utf8mb4 COLLATE utf8mb4_bin DEFAULT NULL,
  `color` varchar(7) CHARACTER SET utf8mb4 COLLATE utf8mb4_bin DEFAULT NULL,
  `note` tinytext CHARACTER SET utf8mb4 COLLATE utf8mb4_bin
) ENGINE=InnoDB DEFAULT CHARSET=utf8mb4;
ALTER TABLE `IFaamjyut` ADD PRIMARY KEY( `id`);
ALTER TABLE `IFaamjyut` ADD UNIQUE( `id`);

INSERT INTO `IFaamjyut` (`id`, `col`, `kind`, `fullname`, `fullname_note`, `color`, `note`) VALUES
"""

# %%
print("1___讀取標籤___")
for n in range(3, len(headers)):
    item = headers[n]
    header_name_value = headers_name[n].value
    header_value      = item.value if item.value else ""
    header_mark_value = int(float(str(headers_mark[n].value))) if headers_mark[n].value else 0
    header_comment    = "" # item.comment.text if item.comment and item.comment.text else ""
    assert header_name_value is None or isinstance(header_name_value, str)
    assert isinstance(header_value, str)
    
    if header_mark_value > 0:
        locale_name_splited: List[str] = (header_name_value+",").split(",")[0:2] if header_name_value else ["", ""]
        header_name = header_value
    else:
        locale_name_splited: List[str] = ["", ""]
        if header_name_value and "," in header_name_value:
            header_name, locale_name_splited[0] = header_name_value.split(",")
        else:
            header_name            = header_value
            locale_name_splited[0] = header_name_value if header_name_value else header_value
            
    # if '"' in header_comment:
    #     header_comment = header_comment.replace('"', "''")
    
    HEADER_INFO_COL_NAME.append(header_name)
    HEADER_INFO_MARK.append(header_mark_value)
    HEADER_INFO_FULL_NAME.append(locale_name_splited)
    HEADER_INFO_NOTE.append(header_comment)
    cell_fill = headers[n].fill
    if len(str(cell_fill.fgColor.rgb))==8:
        cell_color = headers[n].fill.fgColor.rgb[2:]
    else:
        cell_theme = cell_fill.start_color.theme
        cell_tint  = cell_fill.start_color.tint
        cell_color = theme_and_tint_to_rgb(PanCSheet, cell_theme, cell_tint)
    HEADER_INFO_COLOR.append(cell_color)
HEADER_INFO_COL_COUNT = len(headers) # 附註列

# %%
print("2___輸出表頭___")
output_file_name = f"{output_dir}IFaamjyut_{time_str}.sql"
with open(output_file_name, "w", encoding="utf-8") as f:
    f.write(sql_ifaamjyut_header)
    for n in range(HEADER_INFO_COL_COUNT):
        sql_row = "%s(%d, '%s', %d, '%s', '%s', '#%s', '%s')" % (
            ",\n" if n>0 else "",
            n+1, HEADER_INFO_COL_NAME[n], HEADER_INFO_MARK[n], HEADER_INFO_FULL_NAME[n][0], 
            HEADER_INFO_FULL_NAME[n][1], HEADER_INFO_COLOR[n], HEADER_INFO_NOTE[n]
        )
        f.write(sql_row)
    f.write(",\n(%d, '附', 0, '附註', '', '', '');" % (HEADER_INFO_COL_COUNT+1))

sql_row_modal = "(%d, " + ('%s, ' * (HEADER_INFO_COL_COUNT)) + "%s)"
index = 1
col_length = np.ones((HEADER_INFO_COL_COUNT+1, ), dtype="int")
main_sheet_sqls: List[str] = []
dict_n2index: Dict[int, int] = {}
# %%
print("3___解析內容___")
for n in tqdm.trange(7, main_sheet.max_row+1):
    result = format_locale_record(n)
    if not result[2]:
        dict_n2index[n] = -1
        continue
    sql_row_str = sql_row_modal % tuple([index] + result[0])
    col_length = np.max(np.vstack((col_length, result[1])), axis=0)
    main_sheet_sqls.append(sql_row_str)
    dict_n2index[n] = index
    index += 1

# %%
sql_jfaamjyut_header = "CREATE TABLE `JFaamjyut` (\n  `id` int(5) NOT NULL"
for n, i in enumerate(HEADER_INFO_COL_NAME + ["附"]):
    sql_jfaamjyut_header = sql_jfaamjyut_header + \
    ",\n  `%s` varchar(%s) CHARACTER SET utf8mb4 COLLATE utf8mb4_bin DEFAULT NULL" % \
    (i, col_length[n])
sql_jfaamjyut_header = sql_jfaamjyut_header + ",\n  PRIMARY KEY (`id`)\n) ENGINE=InnoDB DEFAULT CHARSET=utf8mb4;\n\n"
sql_jfaamjyut_insert = "INSERT INTO `JFaamjyut` (`" + "`, `".join(["id"] + HEADER_INFO_COL_NAME + ["附"]) + "`) VALUES\n"

print("4___輸出內容___")
output_file_name = f"{output_dir}JFaamjyut_{time_str}_%d.sql"
page = 0
line_per_page = 3000
while page * line_per_page < len(main_sheet_sqls):
    with open(output_file_name % (page+1), "w", encoding="utf-8") as f:
        if page == 0: f.write(sql_jfaamjyut_header)
        f.write(sql_jfaamjyut_insert)
        for n, i in enumerate(main_sheet_sqls[page*line_per_page:(page+1)*line_per_page]):
            if n>0: f.write(",\n")
            f.write(i)
        f.write(";")
    page += 1
